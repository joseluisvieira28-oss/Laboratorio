#!/usr/bin/env python3
"""Bitget public historical depth cadence audit. Source-only; no strategy outcomes."""
from __future__ import annotations
import io,json,statistics,zipfile,time
from datetime import datetime,timezone
from pathlib import Path
import requests
from openpyxl import load_workbook

DATE="2026-09-16"
SYMBOLS=["HOODUSDT","COINUSDT","ARMUSDT","AAPLUSDT"]
API="https://www.bitget.com/v1/statistics/public/download/getPublicDataV2"
HEADERS={
 "User-Agent":"Mozilla/5.0 CryptoLab-BitgetDepthCadence/0.3",
 "Content-Type":"application/json;charset=UTF-8","Accept":"application/json, text/plain, */*",
 "Origin":"https://www.bitget.com","Referer":"https://www.bitget.com/data-download",
 "terminalType":"1","locale":"en_US","language":"en_US","securityNew":"true"
}
OUT=Path("artifacts/cross_venue_stock/bitget_depth_cadence_v03")

def to_ms(v):
    if v is None:return None
    if isinstance(v,datetime):
        dt=v if v.tzinfo else v.replace(tzinfo=timezone.utc)
        return int(dt.timestamp()*1000)
    if isinstance(v,(int,float)):
        x=float(v)
        if x>1e14:x/=1000.0
        elif x<1e11:x*=1000.0
        return int(x)
    s=str(v).strip()
    if not s:return None
    try:
        x=float(s)
        return to_ms(x)
    except:pass
    try:
        dt=datetime.fromisoformat(s.replace("Z","+00:00"))
        if dt.tzinfo is None:dt=dt.replace(tzinfo=timezone.utc)
        return int(dt.timestamp()*1000)
    except:return None

def get_rows(sym,dept):
    payload={"displaySymbol":[sym],"businessLine":2,"businessType":3,"dateType":1,
             "beginTimeStr":DATE,"endTimeStr":DATE,"deptType":dept}
    r=None
    for attempt in range(6):
        r=requests.post(API,json=payload,headers=HEADERS,timeout=60)
        if r.status_code==200:
            break
        if r.status_code==429:
            retry=r.headers.get("Retry-After")
            try:delay=max(1.0,float(retry)) if retry else 2.0*(attempt+1)
            except Exception:delay=2.0*(attempt+1)
            time.sleep(min(delay,12.0))
            continue
        r.raise_for_status()
    if r is None or r.status_code!=200:
        raise RuntimeError(f"DEPTH_API_HTTP_{None if r is None else r.status_code}_{sym}_{dept}")
    j=r.json();data=j.get("data") or []
    if not data or not data[0].get("fileUrl"):raise RuntimeError(f"NO_FILE_{sym}_{dept}")
    q=requests.get(data[0]["fileUrl"],headers={"User-Agent":HEADERS["User-Agent"]},timeout=120);q.raise_for_status()
    z=zipfile.ZipFile(io.BytesIO(q.content));names=z.namelist()
    if len(names)!=1:raise RuntimeError("ZIP_IDENTITY")
    wb=load_workbook(io.BytesIO(z.read(names[0])),read_only=True,data_only=True)
    ws=wb.active;it=ws.iter_rows(values_only=True);hdr=[str(x).strip() for x in next(it)]
    ts=[]
    for row in it:
        if not row:continue
        t=to_ms(row[0])
        if t is not None:ts.append(t)
    ts=sorted(set(ts))
    return hdr,ts

def pct(xs,p):
    if not xs:return None
    ys=sorted(xs);i=min(len(ys)-1,max(0,int(round((len(ys)-1)*p))))
    return ys[i]

def window_coverage(ts,max_age_ms):
    a=int(datetime.fromisoformat(DATE+"T14:31:00+00:00").timestamp()*1000)
    b=int(datetime.fromisoformat(DATE+"T18:44:59+00:00").timestamp()*1000)
    # Evaluate every whole second as a neutral source-quality grid.
    j=0;last=None;ok=0;n=0;ages=[]
    for t in range(a,b+1,1000):
        while j<len(ts) and ts[j]<=t:
            last=ts[j];j+=1
        n+=1
        if last is not None:
            age=t-last
            if age>=0:
                ages.append(age)
                if age<=max_age_ms:ok+=1
    return {
      "grid_seconds":n,
      "observable_seconds":ok,
      "observable_fraction_at_5s":ok/n if n else None,
      "median_latest_snapshot_age_ms":statistics.median(ages) if ages else None,
      "p90_latest_snapshot_age_ms":pct(ages,.90) if ages else None
    }

def main():
    OUT.mkdir(parents=True,exist_ok=True)
    report={"audit_id":"BITGET_DEPTH_CADENCE_SOURCE_AUDIT_V0.3","source_only":True,
            "burned_date":DATE,"max_quote_staleness_seconds_frozen":5,"symbols":{}}
    for sym in SYMBOLS:
        report["symbols"][sym]={}
        for dept in (1,2):
            hdr,ts=get_rows(sym,dept)
            diffs=[ts[i]-ts[i-1] for i in range(1,len(ts)) if ts[i]>ts[i-1]]
            rec={
              "dept_type":dept,"header":hdr,"unique_snapshots":len(ts),
              "first_timestamp_ms":ts[0] if ts else None,"last_timestamp_ms":ts[-1] if ts else None,
              "median_gap_ms":statistics.median(diffs) if diffs else None,
              "p90_gap_ms":pct(diffs,.90) if diffs else None,
              "p99_gap_ms":pct(diffs,.99) if diffs else None,
              "max_gap_ms":max(diffs) if diffs else None,
              "share_gaps_le_1000ms":sum(x<=1000 for x in diffs)/len(diffs) if diffs else None,
              "share_gaps_le_5000ms":sum(x<=5000 for x in diffs)/len(diffs) if diffs else None,
              "signal_window_neutral_coverage":window_coverage(ts,5000)
            }
            report["symbols"][sym][str(dept)]=rec
            print(sym,"dept",dept,"snapshots",len(ts),"median_gap_ms",rec["median_gap_ms"],
                  "coverage_5s",rec["signal_window_neutral_coverage"]["observable_fraction_at_5s"])
            time.sleep(1.0)
    l1=[report["symbols"][s]["1"]["signal_window_neutral_coverage"]["observable_fraction_at_5s"] for s in SYMBOLS]
    l500=[report["symbols"][s]["2"]["signal_window_neutral_coverage"]["observable_fraction_at_5s"] for s in SYMBOLS]
    report["level1_mean_neutral_5s_coverage"]=sum(l1)/len(l1)
    report["level500_mean_neutral_5s_coverage"]=sum(l500)/len(l500)
    if max(l1)<0.5:
        report["verdict"]="LEVEL1_HISTORICAL_CADENCE_STRUCTURALLY_INCOMPATIBLE_WITH_FROZEN_90PCT_OBSERVABILITY_GATE"
    elif min(l1)>=0.9:
        report["verdict"]="LEVEL1_CADENCE_COMPATIBLE__SIGNAL_SPECIFIC_GAPS_REQUIRE_AUDIT"
    else:
        report["verdict"]="LEVEL1_CADENCE_PARTIAL__FURTHER_SOURCE_DIAGNOSTIC_REQUIRED"
    report.update({"strategy_outcomes_opened":0,"private_endpoints_used":False,"account_reads":False,
                   "orders":False,"exchange_mutation":False,"live_trading_authorized":False})
    (OUT/"BITGET_DEPTH_CADENCE_SOURCE_AUDIT_V03.json").write_text(json.dumps(report,indent=2,sort_keys=True))
    print(json.dumps({"verdict":report["verdict"],
      "level1_mean_neutral_5s_coverage":report["level1_mean_neutral_5s_coverage"],
      "level500_mean_neutral_5s_coverage":report["level500_mean_neutral_5s_coverage"]},indent=2))
if __name__=="__main__":main()
