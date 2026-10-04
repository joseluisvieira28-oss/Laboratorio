#!/usr/bin/env python3
from __future__ import annotations
import argparse,bisect,csv,hashlib,io,json,math,statistics,time,zipfile
from datetime import date,datetime,timedelta,timezone
from decimal import Decimal
from pathlib import Path
import requests
from openpyxl import load_workbook

HERE=Path(__file__).resolve().parent
RULE=json.loads((HERE/"BITGET_MAKER_MICROSTRUCTURE_RULE_V0.2.json").read_text())
IDENT=json.loads((HERE/"BITGET_MAKER_PARENT_SIGNAL_IDENTITY_V0.2.json").read_text())
OUT=Path("artifacts/cross_venue_stock/bitget_maker_v02/assets")
UA="CryptoLab-Bitget-MakerV02/0.2"
DEPTH_API="https://www.bitget.com/v1/statistics/public/download/getPublicDataV2"
DEPTH_HEADERS={
 "User-Agent":"Mozilla/5.0 "+UA,
 "Content-Type":"application/json;charset=UTF-8","Accept":"application/json, text/plain, */*",
 "Origin":"https://www.bitget.com","Referer":"https://www.bitget.com/data-download",
 "terminalType":"1","locale":"en_US","language":"en_US","securityNew":"true"
}

def sha(b):return hashlib.sha256(b).hexdigest()
def mean(xs):return sum(xs)/len(xs) if xs else None
def med(xs):return statistics.median(xs) if xs else None
def sgn(x):return 1 if x>0 else (-1 if x<0 else 0)
def sec(d,hh,mm):return int(datetime(d.year,d.month,d.day,hh,mm,tzinfo=timezone.utc).timestamp())
def req(method,url,**kwargs):
    last=None
    for i in range(5):
        try:
            r=requests.request(method,url,timeout=kwargs.pop("timeout",80),**kwargs)
            if r.status_code!=200:raise RuntimeError(f"HTTP_{r.status_code}:{url}")
            return r
        except Exception as e:
            last=e;time.sleep(.6*(i+1))
    raise last

def sessions():
    a=date.fromisoformat(RULE["parent_signal_rule"]["window_start"])
    b=date.fromisoformat(RULE["parent_signal_rule"]["window_end_inclusive"])
    out=[];d=a
    while d<=b:
        if d.weekday()<5:out.append(d)
        d+=timedelta(days=1)
    return out

def fetch_binance(symbol,d):
    day=d.isoformat()
    r=req("GET",f"https://data.binance.vision/data/futures/um/daily/klines/{symbol}/1m/{symbol}-1m-{day}.zip",
          headers={"User-Agent":UA},timeout=100)
    z=zipfile.ZipFile(io.BytesIO(r.content));names=z.namelist()
    if len(names)!=1:raise RuntimeError("BINANCE_ZIP_IDENTITY")
    lo=sec(d,14,29);hi=sec(d,18,59);out={}
    for row in csv.reader(io.TextIOWrapper(z.open(names[0]),encoding="utf-8")):
        try:t=int(row[0])//1000;p=float(row[4])
        except:continue
        if lo<=t<=hi:out[t+60]=p
    return out,sha(r.content)

def fetch_bitget_candles(symbol,d):
    out={};hs=[]
    for a,b in [(sec(d,14,29),sec(d,16,44)),(sec(d,16,45),sec(d,18,59))]:
        r=req("GET","https://api.bitget.com/api/v2/mix/market/history-candles",
          params={"symbol":symbol,"productType":"USDT-FUTURES","granularity":"1m",
                  "startTime":str(a*1000),"endTime":str(b*1000),"limit":"200"},
          headers={"User-Agent":UA})
        hs.append(sha(r.content));j=r.json()
        if j.get("code")!="00000":raise RuntimeError(f"BITGET_KLINE_CODE_{j.get('code')}")
        for row in j.get("data") or []:
            try:out[int(row[0])//1000+60]=float(row[4])
            except:pass
    return out,hs

def reconstruct_signals(bn,bg,d):
    R=RULE["parent_signal_rule"]
    start=sec(d,14,31);stop=sec(d,18,44);next_allowed=start;out=[]
    for t in sorted(set(bn)&set(bg)):
        if t<start or t>stop or t<next_allowed:continue
        prev=t-60
        if prev not in bn or prev not in bg:continue
        lr=10000*(bn[t]/bn[prev]-1)
        tr=10000*(bg[t]/bg[prev]-1)
        gap=lr-tr
        if abs(lr)<R["leader_shock_threshold_bps"] or sgn(gap)!=sgn(lr) or abs(gap)<R["lag_gap_threshold_bps"]:
            continue
        out.append({"signal_ts":t,"direction":sgn(lr),"leader_return_bps":lr,"lag_gap_bps":gap})
        next_allowed=t+R["cooldown_min"]*60
    return out

def depth_file(symbol,d):
    ds=d.isoformat()
    payload={"displaySymbol":[symbol],"businessLine":2,"businessType":3,"dateType":1,
             "beginTimeStr":ds,"endTimeStr":ds,"deptType":1}
    r=req("POST",DEPTH_API,json=payload,headers=DEPTH_HEADERS,timeout=60)
    j=r.json()
    if str(j.get("code")) not in ("200","00000"):raise RuntimeError(f"DEPTH_API_CODE_{j.get('code')}")
    data=j.get("data") or []
    if not data or not data[0].get("fileUrl"):raise RuntimeError(f"DEPTH_NO_FILE_{symbol}_{ds}")
    u=data[0]["fileUrl"]
    q=req("GET",u,headers={"User-Agent":UA,"Referer":"https://www.bitget.com/data-download"},timeout=120)
    z=zipfile.ZipFile(io.BytesIO(q.content));names=z.namelist()
    if len(names)!=1 or not names[0].lower().endswith(".xlsx"):raise RuntimeError("DEPTH_XLSX_IDENTITY")
    xlsx=z.read(names[0])
    wb=load_workbook(io.BytesIO(xlsx),read_only=True,data_only=True)
    ws=wb.active;it=ws.iter_rows(values_only=True);hdr=[str(x).strip() for x in next(it)]
    expected=["timestamp","ask_price","bid_price","ask_volume","bid_volume"]
    if hdr!=expected:raise RuntimeError(f"DEPTH_SCHEMA_{hdr}")
    rows=[]
    for row in it:
        if not row or row[0] is None:continue
        try:
            rows.append((int(row[0]),Decimal(str(row[1])),Decimal(str(row[2])),
                         Decimal(str(row[3])),Decimal(str(row[4]))))
        except Exception:continue
    rows.sort(key=lambda x:x[0])
    return rows,{"zip_sha256":sha(q.content),"xlsx_sha256":sha(xlsx),"rows":len(rows),"file_url":u}

def fetch_trades_chunk(symbol,start_ms,end_ms,depth=0):
    r=req("GET","https://api.bitget.com/api/v2/mix/market/fills-history",
      params={"symbol":symbol,"productType":"USDT-FUTURES","startTime":str(start_ms),
              "endTime":str(end_ms),"limit":"1000"},
      headers={"User-Agent":UA})
    j=r.json()
    if j.get("code")!="00000":raise RuntimeError(f"FILLS_CODE_{j.get('code')}")
    rows=j.get("data") or []
    digest=sha(r.content)
    if len(rows)>=1000 and end_ms-start_ms>1500 and depth<20:
        mid=(start_ms+end_ms)//2
        a,ha=fetch_trades_chunk(symbol,start_ms,mid,depth+1)
        b,hb=fetch_trades_chunk(symbol,mid+1,end_ms,depth+1)
        return a+b,[digest]+ha+hb
    out=[]
    for x in rows:
        try:out.append((int(x["ts"]),Decimal(str(x["price"])),Decimal(str(x["size"])),str(x.get("tradeId",""))))
        except:pass
    return out,[digest]

def fetch_trades_day(symbol,d):
    start=sec(d,14,25)*1000;end=sec(d,19,2)*1000
    rows,hs=fetch_trades_chunk(symbol,start,end)
    uniq={}
    for x in rows:uniq[(x[0],x[3],x[1],x[2])]=x
    out=sorted(uniq.values(),key=lambda x:x[0])
    return out,hs

def latest_snapshot(rows,t,max_age):
    ts=[x[0] for x in rows]
    i=bisect.bisect_right(ts,t)-1
    if i<0:return None
    x=rows[i]
    if t-x[0]>max_age:return None
    return x

def trades_between(trades,a_sec,b_sec):
    lo=bisect.bisect_left([x[0] for x in trades],a_sec*1000)
    hi=bisect.bisect_right([x[0] for x in trades],b_sec*1000+999)
    return trades[lo:hi]

def maker_fill(trades,start_sec,ttl,side,price,queue):
    cum=Decimal("0")
    for ts_ms,p,size,_ in trades_between(trades,start_sec,start_sec+ttl):
        if side=="BUY":
            if p<price:return {"filled":True,"fill_ts_ms":ts_ms,"price":price,"reason":"PRICE_THROUGH"}
            if p==price:
                cum+=size
                if cum>=queue:return {"filled":True,"fill_ts_ms":ts_ms,"price":price,"reason":"QUEUE_CLEAR"}
        else:
            if p>price:return {"filled":True,"fill_ts_ms":ts_ms,"price":price,"reason":"PRICE_THROUGH"}
            if p==price:
                cum+=size
                if cum>=queue:return {"filled":True,"fill_ts_ms":ts_ms,"price":price,"reason":"QUEUE_CLEAR"}
    return {"filled":False,"consumed_at_price":str(cum)}

def thirds(vals):
    n=len(vals)
    if not n:return [None,None,None]
    c=[0,n//3,(2*n)//3,n]
    return [mean(vals[c[i]:c[i+1]]) for i in range(3)]

def binom_tail(w,n):
    return sum(math.comb(n,k) for k in range(w,n+1))/(2**n) if n else None

def main():
    ap=argparse.ArgumentParser();ap.add_argument("--symbol",required=True);args=ap.parse_args()
    symbol=args.symbol
    if symbol not in RULE["candidates"]:raise SystemExit("FAIL_CLOSED_UNKNOWN_SYMBOL")
    E=RULE["passive_execution_model"];ss=sessions()
    all_signals=[];data={};hashes={}
    for d in ss:
        ds=d.isoformat()
        bn,bnh=fetch_binance(symbol,d)
        bg,bgh=fetch_bitget_candles(symbol,d)
        sigs=reconstruct_signals(bn,bg,d)
        depth,dh=depth_file(symbol,d)
        trades,ths=fetch_trades_day(symbol,d)
        all_signals.extend([{**s,"date":ds} for s in sigs])
        data[ds]={"depth":depth,"trades":trades}
        hashes[ds]={"binance":bnh,"bitget_candles":bgh,"depth":dh,"trade_response_sha256":ths,
                    "signal_count":len(sigs),"trade_count":len(trades)}
        time.sleep(.03)
    expected=IDENT["expected_signal_counts"][symbol]
    if len(all_signals)!=expected:
        raise SystemExit(f"FAIL_CLOSED_PARENT_SIGNAL_IDENTITY expected={expected} got={len(all_signals)}")
    obs=0;entries=0;entered_sessions=set();closed=[];unverifiable=0;maker_exits=0;forced_exits=0
    entry_fill_reasons={};exit_fill_reasons={}
    for s in all_signals:
        ds=s["date"];depth=data[ds]["depth"];trades=data[ds]["trades"];t=s["signal_ts"];direction=s["direction"]
        place_t=t+E["strategy_latency_seconds"]
        snap=latest_snapshot(depth,place_t,E["max_quote_staleness_seconds"])
        if snap is None:
            unverifiable+=1;continue
        obs+=1
        _,ask,bid,askv,bidv=snap
        if direction>0:
            entry_side="BUY";entry_price=bid;queue=bidv
        else:
            entry_side="SELL";entry_price=ask;queue=askv
        ef=maker_fill(trades,place_t,E["entry_ttl_seconds"],entry_side,entry_price,queue)
        if not ef["filled"]:continue
        entries+=1;entered_sessions.add(ds)
        entry_fill_reasons[ef["reason"]]=entry_fill_reasons.get(ef["reason"],0)+1
        exit_place=t+60+E["exit_strategy_latency_seconds"]
        xsnap=latest_snapshot(depth,exit_place,E["max_quote_staleness_seconds"])
        if xsnap is None:
            unverifiable+=1
            closed.append({"data_fail":True,"date":ds,"signal_ts":t});continue
        _,xask,xbid,xaskv,xbidv=xsnap
        if direction>0:
            exit_side="SELL";exit_price=xask;xqueue=xaskv
        else:
            exit_side="BUY";exit_price=xbid;xqueue=xbidv
        xf=maker_fill(trades,exit_place,E["exit_ttl_seconds"],exit_side,exit_price,xqueue)
        if xf["filled"]:
            maker_exits+=1;fee=Decimal(str(E["fees_bps"]["maker_maker_roundtrip"]))
            px=xf["price"];exit_kind="MAKER";reason=xf["reason"]
            exit_fill_reasons[reason]=exit_fill_reasons.get(reason,0)+1
        else:
            forced_t=exit_place+E["exit_ttl_seconds"]
            fsnap=latest_snapshot(depth,forced_t,E["forced_exit_quote_max_staleness_seconds"])
            if fsnap is None:
                unverifiable+=1
                closed.append({"data_fail":True,"date":ds,"signal_ts":t});continue
            _,fask,fbid,_,_=fsnap
            px=fbid if direction>0 else fask
            fee=Decimal(str(E["fees_bps"]["maker_taker_roundtrip"]))
            forced_exits+=1;exit_kind="FORCED_TAKER";reason="TTL"
        gross=Decimal(direction)*Decimal("10000")*(px/entry_price-Decimal("1"))
        net=gross-fee
        closed.append({"data_fail":False,"date":ds,"signal_ts":t,"direction":direction,
          "entry_price":str(entry_price),"exit_price":str(px),"exit_kind":exit_kind,
          "gross_bps":float(gross),"fee_bps":float(fee),"net_bps":float(net),
          "entry_fill_reason":ef["reason"],"exit_fill_reason":reason})
    data_fail_entered=sum(1 for x in closed if x.get("data_fail"))
    executed=[x for x in closed if not x.get("data_fail")]
    vals=[x["net_bps"] for x in executed];wins=sum(x>0 for x in vals)
    distinct_exec=len({x["date"] for x in executed})
    coverage=obs/len(all_signals) if all_signals else 0
    entered_complete=(data_fail_entered==0)
    result={
      "symbol":symbol,"parent_signal_count":len(all_signals),"expected_parent_signal_count":expected,
      "placement_observable_count":obs,"placement_observability":coverage,
      "data_unverifiable_count":unverifiable,"entered_positions":entries,
      "entered_distinct_sessions":len(entered_sessions),"closed_positions":len(executed),
      "closed_distinct_sessions":distinct_exec,"entered_positions_with_data_failure":data_fail_entered,
      "maker_exit_count":maker_exits,"forced_taker_exit_count":forced_exits,
      "entry_fill_rate":entries/obs if obs else 0,
      "maker_exit_rate_among_closed":maker_exits/len(executed) if executed else 0,
      "forced_taker_exit_rate_among_closed":forced_exits/len(executed) if executed else 0,
      "mean_realized_net_bps":mean(vals),"median_realized_net_bps":med(vals),
      "wins":wins,"win_rate":wins/len(vals) if vals else None,
      "chronological_third_mean_net_bps":thirds(vals),
      "p_value_one_sided_binomial_vs_50":binom_tail(wins,len(vals)),
      "mean_net_bps_per_parent_signal_with_unfilled_zero":sum(vals)/len(all_signals) if all_signals else None,
      "entry_fill_reasons":entry_fill_reasons,"exit_fill_reasons":exit_fill_reasons,
      "all_entered_positions_closed_or_data_fail":entered_complete,
      "executions":executed
    }
    G=RULE["execution_gate"]
    pre=bool(coverage>=G["min_parent_signal_quote_observability"] and
             entries>=G["min_entered_positions"] and len(entered_sessions)>=G["min_distinct_entered_sessions"] and
             entered_complete and vals and mean(vals)>G["mean_realized_net_bps_gt"] and
             med(vals)>G["median_realized_net_bps_gt"] and wins/len(vals)>G["win_rate_gt"] and
             all(x is not None and x>G["all_chronological_thirds_mean_net_bps_gt"] for x in thirds(vals)))
    receipt={"family_id":RULE["family_id"],"symbol":symbol,"pre_holm_eligible":pre,
      "result":result,"source_hashes":hashes,"post_outcome_tuning":False,
      "private_endpoints_used":False,"account_reads":False,"orders":False,"exchange_mutation":False,
      "live_trading_authorized":False}
    OUT.mkdir(parents=True,exist_ok=True)
    (OUT/(symbol+".json")).write_text(json.dumps(receipt,indent=2,sort_keys=True))
    print(json.dumps({"symbol":symbol,"pre_holm_eligible":pre,
      "signals":len(all_signals),"observable":obs,"entries":entries,"closed":len(executed),
      "maker_exits":maker_exits,"forced_exits":forced_exits,
      "mean_net_bps":result["mean_realized_net_bps"],"median_net_bps":result["median_realized_net_bps"],
      "win_rate":result["win_rate"],"thirds":result["chronological_third_mean_net_bps"],
      "p":result["p_value_one_sided_binomial_vs_50"],
      "per_parent_signal_bps":result["mean_net_bps_per_parent_signal_with_unfilled_zero"]},indent=2))

if __name__=="__main__":main()
