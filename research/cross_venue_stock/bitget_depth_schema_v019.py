#!/usr/bin/env python3
"""Strict source-schema validation for Bitget historical depth XLSX archives. No market values persisted."""
from __future__ import annotations
import io,json,hashlib,zipfile
from datetime import datetime
from pathlib import Path
import requests
from openpyxl import load_workbook

DATE="2026-09-16"
SYMBOLS=["HOODUSDT","COINUSDT","ARMUSDT","AAPLUSDT"]
API="https://www.bitget.com/v1/statistics/public/download/getPublicDataV2"
HEADERS={
 "User-Agent":"Mozilla/5.0 CryptoLab-BitgetDepthSchema/0.1.9",
 "Content-Type":"application/json;charset=UTF-8","Accept":"application/json, text/plain, */*",
 "Origin":"https://www.bitget.com","Referer":"https://www.bitget.com/data-download",
 "terminalType":"1","locale":"en_US","language":"en_US","securityNew":"true"
}
OUT=Path("artifacts/cross_venue_stock/bitget_depth_schema_v019")

def sha(b):return hashlib.sha256(b).hexdigest()

def get_file(sym,dept):
    payload={"displaySymbol":[sym],"businessLine":2,"businessType":3,"dateType":1,
             "beginTimeStr":DATE,"endTimeStr":DATE,"deptType":dept}
    r=requests.post(API,json=payload,headers=HEADERS,timeout=60);r.raise_for_status()
    j=r.json();data=j.get("data") or []
    if not isinstance(data,list) or not data or not data[0].get("fileUrl"):
        raise RuntimeError(f"NO_FILE_{sym}_{dept}:{j}")
    meta={k:data[0].get(k) for k in ("fileName","fileUrl","dateTime","dateTimeStr")}
    q=requests.get(meta["fileUrl"],headers={"User-Agent":HEADERS["User-Agent"]},timeout=120);q.raise_for_status()
    z=zipfile.ZipFile(io.BytesIO(q.content));names=z.namelist()
    if len(names)!=1 or not names[0].lower().endswith(".xlsx"):
        raise RuntimeError(f"XLSX_IDENTITY_{sym}_{dept}_{names}")
    xlsx=z.read(names[0])
    return meta,q.content,names[0],xlsx

def norm_header(v):
    return "" if v is None else str(v).strip()

def parse_ts(v):
    if v is None:return None
    if isinstance(v,datetime):return v.isoformat()
    if isinstance(v,(int,float)):
        # preserve numeric timestamp only; do not infer price-like columns.
        return str(v)
    s=str(v).strip()
    return s[:80] if s else None

def schema(xlsx):
    wb=load_workbook(io.BytesIO(xlsx),read_only=True,data_only=True)
    sheets=[]
    for ws in wb.worksheets:
        it=ws.iter_rows(values_only=True)
        first=next(it,None)
        header=[norm_header(x) for x in first] if first else []
        rows=0;first_ts=None;last_ts=None
        nonnull_counts=[0]*len(header)
        cell_type_sets=[set() for _ in header]
        max_text_len=[0]*len(header)
        for row in it:
            if not any(x is not None for x in row):continue
            rows+=1
            if row:
                tv=parse_ts(row[0])
                if tv is not None:
                    first_ts=tv if first_ts is None else first_ts
                    last_ts=tv
            for i,x in enumerate(row[:len(header)]):
                if x is not None:
                    nonnull_counts[i]+=1
                    cell_type_sets[i].add(type(x).__name__)
                    if isinstance(x,str):max_text_len[i]=max(max_text_len[i],len(x))
        sheets.append({
          "title":ws.title,"data_row_count":rows,"column_count":len(header),"header":header,
          "first_column_first_value":first_ts,"first_column_last_value":last_ts,
          "column_types":[sorted(x) for x in cell_type_sets],
          "nonnull_counts":nonnull_counts,"max_string_lengths":max_text_len
        })
    return sheets

def main():
    OUT.mkdir(parents=True,exist_ok=True)
    report={"gate_id":"BITGET_DEPTH_SCHEMA_V0_1_9_1","source_only":True,"burned_date":DATE,
            "market_depth_prices_persisted":False,"symbols":{}}
    for sym in SYMBOLS:
        report["symbols"][sym]={}
        for dept in (1,2):
            meta,outer,name,xlsx=get_file(sym,dept)
            sch=schema(xlsx)
            report["symbols"][sym][str(dept)]={
              "file_meta":meta,"outer_zip_sha256":sha(outer),"outer_zip_bytes":len(outer),
              "xlsx_member_name":name,"xlsx_sha256":sha(xlsx),"xlsx_bytes":len(xlsx),"sheets":sch
            }
            print(sym,"dept",dept,"file",meta["fileName"])
            for sh in sch:
                print(" SHEET",sh["title"],"rows",sh["data_row_count"],"cols",sh["column_count"],
                      "header",sh["header"],"first",sh["first_column_first_value"],"last",sh["first_column_last_value"],
                      "types",sh["column_types"],"maxlen",sh["max_string_lengths"])
    required_l1={"timestamp","ask_price","bid_price","ask_volume","bid_volume"}
    required_l500={"timestamp","asks","bids"}
    passes={}
    for sym,d in report["symbols"].items():
        passes[sym]={}
        for dept,x in d.items():
            ok=False
            for sh in x["sheets"]:
                hs={h.lower() for h in sh["header"]}
                required = required_l1 if dept=="1" else required_l500
                ok=ok or (required.issubset(hs) and sh["data_row_count"]>0)
            passes[sym][dept]=ok
    report["passes"]=passes
    report["all_level1_schema_pass"]=all(v["1"] for v in passes.values())
    report["all_level500_schema_pass"]=all(v["2"] for v in passes.values())
    report["verdict"]="BITGET_HISTORICAL_DEPTH_SCHEMA_PASS" if report["all_level1_schema_pass"] else "BITGET_HISTORICAL_DEPTH_SCHEMA_BLOCKED"
    report.update({"microstructure_outcomes_opened":0,"private_endpoints_used":False,"account_reads":False,
                   "orders":False,"exchange_mutation":False,"live_trading_authorized":False})
    (OUT/"BITGET_DEPTH_SCHEMA_V019.json").write_text(json.dumps(report,indent=2,sort_keys=True))
    print(json.dumps({"verdict":report["verdict"],"passes":passes},indent=2))

if __name__=="__main__":main()
